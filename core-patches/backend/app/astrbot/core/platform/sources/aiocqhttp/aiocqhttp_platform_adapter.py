import asyncio
import inspect
import io
import itertools
import json
import logging
import re
import time
import urllib.request
import uuid
from collections.abc import Awaitable
from typing import Any, cast

from aiocqhttp import CQHttp, Event
from aiocqhttp.exceptions import ActionFailed

from astrbot.api import logger
from astrbot.api.event import MessageChain
from astrbot.api.message_components import *
from astrbot.api.platform import (
    AstrBotMessage,
    MessageMember,
    MessageType,
    Platform,
    PlatformMetadata,
)
from astrbot.core.platform.astr_message_event import MessageSesion
from astrbot.core.utils.quoted_message.chain_parser import OneBotPayloadParser
from astrbot.core.utils.qq_face import face_name

from ...register import register_platform_adapter
from .aiocqhttp_message_event import *
from .aiocqhttp_message_event import AiocqhttpMessageEvent

# [本地修改] 合并转发里的图片/文件处理上限。
# 深度 0/1/2 表示最多展开三层合并转发。
FORWARD_MAX_DEPTH = 2
FORWARD_MAX_IMAGES = 6
FORWARD_MAX_FILES = 3
FORWARD_MAX_NESTED_FETCH = 5
FORWARD_MAX_TEXT_CHARS = 4000
FILE_MAX_BYTES = 8 * 1024 * 1024
FILE_TEXT_MAX_CHARS = 1500
FILE_PARSE_TIMEOUT_SEC = 25

# 能直接抽成文本的扩展名；其余类型只给文件名。
_FILE_TEXT_EXTS = {
    "txt",
    "md",
    "markdown",
    "json",
    "jsonl",
    "csv",
    "tsv",
    "log",
    "py",
    "java",
    "kt",
    "js",
    "ts",
    "jsx",
    "tsx",
    "html",
    "htm",
    "css",
    "scss",
    "xml",
    "yaml",
    "yml",
    "ini",
    "conf",
    "properties",
    "sql",
    "sh",
    "bat",
    "ps1",
    "c",
    "h",
    "cpp",
    "hpp",
    "cs",
    "go",
    "rs",
    "php",
    "rb",
    "swift",
    "toml",
    "pdf",
    "docx",
    "xlsx",
    "xlsm",
}


def _download_file_bytes(url: str) -> bytes:
    # [本地修改] 只下载 QQ 给出的文件直链，限制大小。
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0"},
    )
    with urllib.request.urlopen(request, timeout=30) as resp:
        data = resp.read(FILE_MAX_BYTES + 1)
    if len(data) > FILE_MAX_BYTES:
        raise ValueError(f"file too large (> {FILE_MAX_BYTES} bytes)")
    return data


def _decode_text_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is not None:
            return str(best)
    except Exception:  # noqa: BLE001
        pass
    return data.decode("utf-8", "ignore")


def _parse_file_bytes(name: str, data: bytes) -> str:
    # [本地修改] 文本/PDF/Word/Excel 转纯文本，依赖都是 AstrBot 自带的。
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in {"pdf"}:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages: list[str] = []
        for index, page in enumerate(reader.pages):
            if index >= 30:
                break
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    if ext in {"docx"}:
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables[:3]:
            for row in table.rows[:20]:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext in {"xlsx", "xlsm"}:
        import openpyxl

        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
        sheet = workbook.active
        rows: list[str] = []
        if sheet is not None:
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 50:
                    break
                cells = ["" if cell is None else str(cell) for cell in row[:20]]
                if any(cell.strip() for cell in cells):
                    rows.append(" | ".join(cells))
        workbook.close()
        return "\n".join(rows)
    return _decode_text_bytes(data)


@register_platform_adapter(
    "aiocqhttp",
    "适用于 OneBot V11 标准的消息平台适配器，支持反向 WebSockets。",
    support_streaming_message=False,
)
class AiocqhttpAdapter(Platform):
    def __init__(
        self,
        platform_config: dict,
        platform_settings: dict,
        event_queue: asyncio.Queue,
    ) -> None:
        super().__init__(platform_config, event_queue)

        self.settings = platform_settings
        self.host = platform_config["ws_reverse_host"]
        self.port = platform_config["ws_reverse_port"]

        self.metadata = PlatformMetadata(
            name="aiocqhttp",
            description="适用于 OneBot 标准的消息平台适配器，支持反向 WebSockets。",
            id=cast(str, self.config.get("id")),
            support_streaming_message=False,
        )

        self.bot = CQHttp(
            use_ws_reverse=True,
            import_name="aiocqhttp",
            api_timeout_sec=180,
            access_token=platform_config.get(
                "ws_reverse_token",
            ),  # 以防旧版本配置不存在
        )

        @self.bot.on_request()
        async def request(event: Event) -> None:
            try:
                abm = await self.convert_message(event)
                if not abm:
                    return
                await self.handle_msg(abm)
            except Exception as e:
                logger.exception(f"Handle request message failed: {e}")
                return

        @self.bot.on_notice()
        async def notice(event: Event) -> None:
            try:
                abm = await self.convert_message(event)
                if abm:
                    await self.handle_msg(abm)
            except Exception as e:
                logger.exception(f"Handle notice message failed: {e}")
                return

        @self.bot.on_message("group")
        async def group(event: Event) -> None:
            try:
                abm = await self.convert_message(event)
                if abm:
                    await self.handle_msg(abm)
            except Exception as e:
                logger.exception(f"Handle group message failed: {e}")
                return

        @self.bot.on_message("private")
        async def private(event: Event) -> None:
            try:
                abm = await self.convert_message(event)
                if abm:
                    await self.handle_msg(abm)
            except Exception as e:
                logger.exception(f"Handle private message failed: {e}")
                return

        @self.bot.on_websocket_connection
        def on_websocket_connection(_) -> None:
            logger.info("aiocqhttp(OneBot v11) 适配器已连接。")

    async def send_by_session(
        self,
        session: MessageSesion,
        message_chain: MessageChain,
    ) -> None:
        is_group = session.message_type == MessageType.GROUP_MESSAGE
        if is_group:
            session_id = session.session_id.split("_")[-1]
        else:
            session_id = session.session_id
        await AiocqhttpMessageEvent.send_message(
            bot=self.bot,
            message_chain=message_chain,
            event=None,  # 这里不需要 event，因为是通过 session 发送的
            is_group=is_group,
            session_id=session_id,
        )
        await super().send_by_session(session, message_chain)

    async def convert_message(self, event: Event) -> AstrBotMessage | None:
        logger.debug(f"[aiocqhttp] RawMessage {event}")

        if event["post_type"] == "message":
            abm = await self._convert_handle_message_event(event)
            if abm.sender.user_id == "2854196310":
                # 屏蔽 QQ 管家的消息
                return None
        elif event["post_type"] == "notice":
            abm = await self._convert_handle_notice_event(event)
        elif event["post_type"] == "request":
            abm = await self._convert_handle_request_event(event)

        return abm

    async def _convert_handle_request_event(self, event: Event) -> AstrBotMessage:
        """OneBot V11 请求类事件"""
        abm = AstrBotMessage()
        abm.self_id = str(event.self_id)
        abm.sender = MessageMember(
            user_id=str(event.user_id), nickname=str(event.user_id)
        )
        abm.type = MessageType.OTHER_MESSAGE
        if event.get("group_id"):
            abm.type = MessageType.GROUP_MESSAGE
            abm.group_id = str(event.group_id)
        else:
            abm.type = MessageType.FRIEND_MESSAGE
        abm.session_id = (
            str(event.group_id)
            if abm.type == MessageType.GROUP_MESSAGE
            else abm.sender.user_id
        )
        abm.message_str = ""
        abm.message = []
        abm.timestamp = int(time.time())
        abm.message_id = uuid.uuid4().hex
        abm.raw_message = event
        return abm

    async def _convert_handle_notice_event(self, event: Event) -> AstrBotMessage:
        """OneBot V11 通知类事件"""
        abm = AstrBotMessage()
        abm.self_id = str(event.self_id)
        abm.sender = MessageMember(
            user_id=str(event.user_id), nickname=str(event.user_id)
        )
        abm.type = MessageType.OTHER_MESSAGE
        if event.get("group_id"):
            abm.group_id = str(event.group_id)
            abm.type = MessageType.GROUP_MESSAGE
        else:
            abm.type = MessageType.FRIEND_MESSAGE
        abm.session_id = (
            str(event.group_id)
            if abm.type == MessageType.GROUP_MESSAGE
            else abm.sender.user_id
        )
        abm.message_str = ""
        abm.message = []
        abm.raw_message = event
        abm.timestamp = int(time.time())
        abm.message_id = uuid.uuid4().hex

        if "sub_type" in event:
            if event["sub_type"] == "poke" and "target_id" in event:
                abm.message.append(Poke(id=str(event["target_id"])))

        return abm

    async def _convert_handle_message_event(
        self,
        event: Event,
        get_reply=True,
    ) -> AstrBotMessage:
        """OneBot V11 消息类事件

        @param event: 事件对象
        @param get_reply: 是否获取回复消息。这个参数是为了防止多个回复嵌套。
        """
        assert event.sender is not None
        abm = AstrBotMessage()
        abm.self_id = str(event.self_id)
        abm.sender = MessageMember(
            str(event.sender["user_id"]),
            event.sender.get("card") or event.sender.get("nickname", "N/A"),
        )
        if event["message_type"] == "group":
            abm.type = MessageType.GROUP_MESSAGE
            abm.group_id = str(event.group_id)
            abm.group = Group(str(event.group_id))
            abm.group.group_name = event.get("group_name")
        elif event["message_type"] == "private":
            abm.type = MessageType.FRIEND_MESSAGE
        abm.session_id = (
            str(event.group_id)
            if abm.type == MessageType.GROUP_MESSAGE
            else abm.sender.user_id
        )

        abm.message_id = str(event.message_id)
        abm.message = []

        message_str = ""
        if not isinstance(event.message, list):
            err = f"aiocqhttp: 无法识别的消息类型: {event.message!s}，此条消息将被忽略。如果您在使用 go-cqhttp，请将其配置文件中的 message.post-format 更改为 array。"
            logger.critical(err)
            try:
                await self.bot.send(event, err)
            except BaseException as e:
                logger.error(f"回复消息失败: {e}")
            raise ValueError(err)

        # 按消息段类型类型适配
        routing_params = {"self_id": event.self_id} if event.self_id else {}
        for t, m_group in itertools.groupby(event.message, key=lambda x: x["type"]):
            a = None
            if t == "text":
                current_text = "".join(m["data"]["text"] for m in m_group).strip()
                if not current_text:
                    # 如果文本段为空，则跳过
                    continue
                message_str += current_text
                a = ComponentTypes[t](text=current_text)
                abm.message.append(a)

            elif t == "file":
                # [本地修改] 直接发来的文件：除了挂 File 组件，还把文件名和
                # 可抽取的正文写进 message_str（file_extract 默认关闭）。
                for m in m_group:
                    file_name = (
                        m["data"].get("file_name", "")
                        or m["data"].get("name", "")
                        or m["data"].get("file", "")
                        or "file"
                    )
                    file_url = ""
                    if m["data"].get("url") and m["data"]["url"].startswith("http"):
                        # Lagrange
                        logger.info("guessing lagrange")
                        file_url = m["data"]["url"]
                    else:
                        try:
                            # Napcat
                            ret = None
                            if abm.type == MessageType.GROUP_MESSAGE:
                                ret = await self.bot.call_action(
                                    action="get_group_file_url",
                                    file_id=m["data"]["file_id"],
                                    group_id=event.group_id,
                                    **routing_params,
                                )
                            elif abm.type == MessageType.FRIEND_MESSAGE:
                                ret = await self.bot.call_action(
                                    action="get_private_file_url",
                                    file_id=m["data"]["file_id"],
                                    **routing_params,
                                )
                            if ret and "url" in ret:
                                file_url = ret["url"]  # https
                                # 优先从 API 返回值获取文件名，其次从原始消息数据获取
                                file_name = (
                                    ret.get("file_name", "")
                                    or ret.get("name", "")
                                    or m["data"].get("file", "")
                                    or m["data"].get("file_name", "")
                                    or "file"
                                )
                            else:
                                logger.error(f"获取文件失败: {ret}")

                        except ActionFailed as e:
                            logger.error(f"获取文件失败: {e}，此消息段将被忽略。")
                        except BaseException as e:
                            logger.error(f"获取文件失败: {e}，此消息段将被忽略。")

                    if file_url:
                        abm.message.append(File(name=file_name, url=file_url))
                    message_str += f"\n[文件:{file_name}]"
                    file_text = await self._extract_file_text(file_name, file_url)
                    if file_text:
                        message_str += (
                            f"\n[文件内容:{file_name}]\n{file_text}\n[/文件内容]"
                        )

            elif t == "reply":
                for m in m_group:
                    if not get_reply:
                        a = ComponentTypes[t](**m["data"])
                        abm.message.append(a)
                    else:
                        try:
                            reply_event_data = await self.bot.call_action(
                                action="get_msg",
                                message_id=int(m["data"]["id"]),
                                **routing_params,
                            )
                            # 添加必要的 post_type 字段，防止 Event.from_payload 报错
                            reply_event_data["post_type"] = "message"
                            new_event = Event.from_payload(reply_event_data)
                            if not new_event:
                                logger.error(
                                    f"无法从回复消息数据构造 Event 对象: {reply_event_data}",
                                )
                                continue
                            abm_reply = await self._convert_handle_message_event(
                                new_event,
                                get_reply=False,
                            )

                            reply_seg = Reply(
                                id=abm_reply.message_id,
                                chain=abm_reply.message,
                                sender_id=abm_reply.sender.user_id,
                                sender_nickname=abm_reply.sender.nickname,
                                time=abm_reply.timestamp,
                                message_str=abm_reply.message_str,
                                text=abm_reply.message_str,  # for compatibility
                                qq=abm_reply.sender.user_id,  # for compatibility
                            )

                            abm.message.append(reply_seg)
                        except BaseException as e:
                            logger.error(f"获取引用消息失败: {e}。")
                            a = ComponentTypes[t](**m["data"])
                            abm.message.append(a)
            elif t == "at":
                first_at_self_processed = False
                # Accumulate @ mention text for efficient concatenation
                at_parts = []

                for m in m_group:
                    try:
                        if m["data"]["qq"] == "all":
                            abm.message.append(At(qq="all", name="全体成员"))
                            continue

                        at_info = await self.bot.call_action(
                            action="get_group_member_info",
                            group_id=event.group_id,
                            user_id=int(m["data"]["qq"]),
                            no_cache=False,
                            **routing_params,
                        )
                        if at_info:
                            nickname = at_info.get("card", "")
                            if nickname == "":
                                at_info = await self.bot.call_action(
                                    action="get_stranger_info",
                                    user_id=int(m["data"]["qq"]),
                                    no_cache=False,
                                    **routing_params,
                                )
                                nickname = at_info.get("nick", "") or at_info.get(
                                    "nickname",
                                    "",
                                )
                            is_at_self = str(m["data"]["qq"]) in {abm.self_id, "all"}

                            abm.message.append(
                                At(
                                    qq=m["data"]["qq"],
                                    name=nickname,
                                ),
                            )

                            if is_at_self and not first_at_self_processed:
                                # 第一个@是机器人，不添加到message_str
                                first_at_self_processed = True
                            else:
                                # 非第一个@机器人或@其他用户，添加到message_str
                                at_parts.append(f" @{nickname}({m['data']['qq']}) ")
                        else:
                            abm.message.append(At(qq=str(m["data"]["qq"]), name=""))
                    except ActionFailed as e:
                        logger.error(f"获取 @ 用户信息失败: {e}，此消息段将被忽略。")
                    except BaseException as e:
                        logger.error(f"获取 @ 用户信息失败: {e}，此消息段将被忽略。")

                message_str += "".join(at_parts)
            elif t == "face":
                # [本地修改] 普通 QQ 表情：LLM 只看 message_str，补上表情名。
                for m in m_group:
                    fid = m["data"].get("id", "")
                    name = face_name(fid)
                    message_str += (
                        f"[QQ表情:{name}]" if name else f"[QQ表情:id={fid}]"
                    )
                    try:
                        abm.message.append(Face(id=int(fid)))
                    except (TypeError, ValueError):
                        pass
            elif t == "mface":
                # [本地修改] 商城/超级表情：用 summary 作为可读文本。
                for m in m_group:
                    label = (
                        m["data"].get("summary")
                        or m["data"].get("emoji_id")
                        or ""
                    )
                    message_str += (
                        f"[动画表情:{label}]" if label else "[动画表情]"
                    )
            elif t == "forward":
                # [本地修改] 合并转发：拉取节点内容并展开成文本，否则 LLM 看不到内容。
                # 同时把节点里的图片/文件补进 abm.message，让模型真的能看到。
                for m in m_group:
                    fid = m["data"].get("id") or m["data"].get("message_id")
                    if not fid:
                        message_str += "[合并转发消息:内容获取失败]"
                        continue
                    media_state = self._new_forward_media_state()
                    text = await self._render_forward_message(
                        fid, routing_params, media_state=media_state
                    )
                    if text:
                        file_blocks = await self._render_forward_file_blocks(
                            media_state["files"]
                        )
                        if file_blocks:
                            text = f"{text}\n" + "\n".join(file_blocks)
                        message_str += f"\n[合并转发消息]\n{text}\n[/合并转发消息]"
                        try:
                            abm.message.append(Forward(id=str(fid)))
                        except Exception:
                            pass
                        for image_url in media_state["images"]:
                            try:
                                abm.message.append(
                                    Image(file=image_url, url=image_url)
                                )
                            except Exception:
                                pass
                        for file_ref in media_state["files"]:
                            try:
                                abm.message.append(
                                    File(
                                        name=file_ref["name"],
                                        url=file_ref["url"],
                                    )
                                )
                            except Exception:
                                pass
                        logger.info(
                            "[本地修改] 合并转发展开：%s 张图片，%s 个文件，%s 字",
                            len(media_state["images"]),
                            len(media_state["files"]),
                            len(text),
                        )
                    else:
                        message_str += "[合并转发消息:内容获取失败]"
            elif t == "markdown":
                for m in m_group:
                    text = m["data"].get("markdown") or m["data"].get("content", "")
                    abm.message.append(Plain(text=text))
                    message_str += text
            else:
                for m in m_group:
                    try:
                        if t not in ComponentTypes:
                            logger.warning(
                                f"不支持的消息段类型，已忽略: {t}, data={m['data']}"
                            )
                            continue
                        a = ComponentTypes[t](**m["data"])
                        abm.message.append(a)
                    except Exception as e:
                        logger.exception(
                            f"消息段解析失败: type={t}, data={m['data']}. {e}"
                        )
                        continue

        abm.timestamp = int(time.time())
        abm.message_str = message_str
        abm.raw_message = event

        return abm

    @staticmethod
    def _new_forward_media_state() -> dict[str, Any]:
        # [本地修改] 一次合并转发的图片/文件收集状态；嵌套层共用。
        return {
            "images": [],
            "image_seen": set(),
            "files": [],
            "file_seen": set(),
        }

    async def _render_forward_message(
        self,
        forward_id: object,
        routing_params: dict,
        *,
        depth: int = 0,
        media_state: dict[str, Any] | None = None,
    ) -> str:
        """[本地修改] 把 OneBot 合并转发消息展开成可读文本。

        同时把节点里的图片 URL、文件信息收集进 media_state，交给调用方
        追加成 Image / File 组件，模型才能真正看到内容。
        """
        if media_state is None:
            media_state = self._new_forward_media_state()
        if depth > FORWARD_MAX_DEPTH:
            return "...[嵌套转发过深]"
        fid = str(forward_id or "").strip()
        if not fid:
            return ""

        payload = None
        for params in ({"message_id": fid}, {"id": fid}):
            try:
                payload = await self.bot.call_action(
                    action="get_forward_msg",
                    **params,
                    **routing_params,
                )
                if isinstance(payload, dict):
                    break
            except ActionFailed as e:
                logger.debug(f"获取合并转发消息失败({params}): {e}")
            except BaseException as e:
                logger.debug(f"获取合并转发消息失败({params}): {e}")
            payload = None
        if not payload:
            return ""

        parsed = OneBotPayloadParser().parse_get_forward_payload(payload)
        parts: list[str] = []
        if parsed["text"]:
            parts.append(
                self._number_forward_images(
                    parsed["text"],
                    parsed["image_refs"],
                    media_state,
                )
            )
        self._collect_forward_files(payload, media_state)
        for nested_id in parsed["forward_ids"][:FORWARD_MAX_NESTED_FETCH]:
            nested = await self._render_forward_message(
                nested_id,
                routing_params,
                depth=depth + 1,
                media_state=media_state,
            )
            if nested:
                parts.append(f"[嵌套转发]\n{nested}")
        text = "\n".join(parts).strip()
        if len(text) > FORWARD_MAX_TEXT_CHARS:
            text = text[:FORWARD_MAX_TEXT_CHARS] + "\n...[内容过长已截断]"
        return text

    @staticmethod
    def _number_forward_images(
        text: str,
        image_refs: list[str],
        media_state: dict[str, Any],
    ) -> str:
        """把 [Image] 占位符换成像 [图片1] 这样带序号的文本。

        编号和实际追加的 Image 组件顺序一致；超出上限或没有 URL 的
        只保留 [图片]，避免模型凭空引用不存在的图。
        """
        ref_iter = iter(image_refs)
        out: list[str] = []
        for chunk in re.split(r"(\[Image\])", text):
            if chunk != "[Image]":
                out.append(chunk)
                continue
            ref = next(ref_iter, None)
            label = AiocqhttpAdapter._register_forward_image(ref, media_state)
            out.append(label)
        # 有些图片是以 file 段发进来的，正文里没有 [Image] 占位符。
        for ref in ref_iter:
            label = AiocqhttpAdapter._register_forward_image(ref, media_state)
            if label and label != "[图片]":
                out.append(f" {label}")
        return "".join(out)

    @staticmethod
    def _register_forward_image(
        ref: str | None,
        media_state: dict[str, Any],
    ) -> str:
        if not isinstance(ref, str) or not ref.strip():
            return "[图片]"
        ref = ref.strip()
        if ref in media_state["image_seen"]:
            return "[图片]"
        if len(media_state["images"]) >= FORWARD_MAX_IMAGES:
            media_state["image_seen"].add(ref)
            return "[图片]"
        media_state["image_seen"].add(ref)
        media_state["images"].append(ref)
        return f"[图片{len(media_state['images'])}]"

    @staticmethod
    def _collect_forward_files(
        payload: dict[str, Any],
        media_state: dict[str, Any],
        *,
        depth: int = 0,
    ) -> None:
        """收集合并转发节点里的 file 段（含内联嵌套节点）。"""
        data = payload.get("data") if isinstance(payload, dict) else None
        if not isinstance(data, dict):
            data = payload
        nodes = (
            data.get("messages")
            or data.get("message")
            or data.get("nodes")
            or data.get("nodeList")
            or data.get("content")
        )
        if not isinstance(nodes, list) or depth > FORWARD_MAX_DEPTH:
            return

        for node in nodes:
            if not isinstance(node, dict):
                continue
            chain = node.get("message") or node.get("content") or []
            if isinstance(chain, str):
                try:
                    parsed_chain = json.loads(chain)
                except Exception:
                    parsed_chain = None
                chain = parsed_chain if isinstance(parsed_chain, list) else []
            if not isinstance(chain, list):
                continue
            for seg in chain:
                if not isinstance(seg, dict):
                    continue
                seg_type = seg.get("type")
                seg_data = seg.get("data") if isinstance(seg.get("data"), dict) else {}
                if seg_type == "file":
                    name = (
                        seg_data.get("name")
                        or seg_data.get("file_name")
                        or seg_data.get("file")
                        or "file"
                    )
                    url = seg_data.get("url") or ""
                    if not str(url).startswith("http"):
                        raw_file = seg_data.get("file") or ""
                        url = raw_file if str(raw_file).startswith("http") else ""
                    key = str(url or name)
                    if key in media_state["file_seen"]:
                        continue
                    if len(media_state["files"]) >= FORWARD_MAX_FILES:
                        continue
                    media_state["file_seen"].add(key)
                    media_state["files"].append(
                        {"name": str(name), "url": str(url or "")}
                    )
                elif seg_type in ("forward", "forward_msg", "nodes"):
                    AiocqhttpAdapter._collect_forward_files(
                        {"data": seg_data},
                        media_state,
                        depth=depth + 1,
                    )

    async def _render_forward_file_blocks(
        self, files: list[dict[str, str]]
    ) -> list[str]:
        blocks: list[str] = []
        for file_ref in files:
            content = await self._extract_file_text(
                file_ref["name"], file_ref["url"]
            )
            if content:
                blocks.append(
                    f"[文件内容:{file_ref['name']}]\n{content}\n[/文件内容]"
                )
        return blocks

    async def _extract_file_text(self, name: str, url: str) -> str:
        """[本地修改] 下载并抽取文本类文件的正文；失败就静默跳过。"""
        if not url or not name:
            return ""
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if ext not in _FILE_TEXT_EXTS:
            return ""
        try:
            data = await asyncio.wait_for(
                asyncio.to_thread(_download_file_bytes, url),
                timeout=FILE_PARSE_TIMEOUT_SEC + 10,
            )
        except Exception as e:  # noqa: BLE001
            logger.debug("[本地修改] 文件下载失败 %s: %s", name, e)
            return ""
        if not data:
            return ""
        try:
            text = await asyncio.wait_for(
                asyncio.to_thread(_parse_file_bytes, name, data),
                timeout=FILE_PARSE_TIMEOUT_SEC,
            )
        except Exception as e:  # noqa: BLE001
            logger.debug("[本地修改] 文件解析失败 %s: %s", name, e)
            return ""
        text = (text or "").strip()
        if not text:
            return f"（{ext} 文件没有可抽取的文字）"
        if len(text) > FILE_TEXT_MAX_CHARS:
            text = text[:FILE_TEXT_MAX_CHARS] + "\n...[文件内容过长已截断]"
        return text

    def run(self) -> Awaitable[Any]:
        if not self.host or not self.port:
            logger.warning(
                "aiocqhttp: 未配置 ws_reverse_host 或 ws_reverse_port，将使用默认值：http://0.0.0.0:6199",
            )
            self.host = "0.0.0.0"
            self.port = 6199

        coro = self.bot.run_task(
            host=self.host,
            port=int(self.port),
            shutdown_trigger=self.shutdown_trigger_placeholder,
        )

        for handler in logging.root.handlers[:]:
            logging.root.removeHandler(handler)
        logging.getLogger("aiocqhttp").setLevel(logging.ERROR)
        self.shutdown_event = asyncio.Event()
        return coro

    async def terminate(self) -> None:
        if hasattr(self, "shutdown_event"):
            self.shutdown_event.set()
        await self._close_reverse_ws_connections()

    async def _close_reverse_ws_connections(self) -> None:
        api_clients = getattr(self.bot, "_wsr_api_clients", None)
        event_clients = getattr(self.bot, "_wsr_event_clients", None)

        ws_clients: set[Any] = set()
        if isinstance(api_clients, dict):
            ws_clients.update(api_clients.values())
        if isinstance(event_clients, set):
            ws_clients.update(event_clients)

        close_tasks: list[Awaitable[Any]] = []
        for ws in ws_clients:
            close_func = getattr(ws, "close", None)
            if not callable(close_func):
                continue
            try:
                close_result = close_func(code=1000, reason="Adapter shutdown")
            except TypeError:
                close_result = close_func()
            except Exception:
                continue

            if inspect.isawaitable(close_result):
                close_tasks.append(close_result)

        if close_tasks:
            await asyncio.gather(*close_tasks, return_exceptions=True)

        if isinstance(api_clients, dict):
            api_clients.clear()
        if isinstance(event_clients, set):
            event_clients.clear()

    async def shutdown_trigger_placeholder(self) -> None:
        await self.shutdown_event.wait()
        logger.info("aiocqhttp 适配器已被关闭")

    def meta(self) -> PlatformMetadata:
        return self.metadata

    def create_event(self, message: AstrBotMessage) -> AiocqhttpMessageEvent:
        """Creates an aiocqhttp message event.

        Args:
            message: AstrBot message object to wrap.

        Returns:
            Created aiocqhttp message event.
        """
        return AiocqhttpMessageEvent(
            message_str=message.message_str,
            message_obj=message,
            platform_meta=self.meta(),
            session_id=message.session_id,
            bot=self.bot,
        )

    async def handle_msg(self, message: AstrBotMessage) -> None:
        self.commit_event(self.create_event(message))

    def get_client(self) -> CQHttp:
        return self.bot
